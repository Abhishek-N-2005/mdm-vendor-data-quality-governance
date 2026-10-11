"""Build the Excel workbook: raw + golden vendor data with LIVE formula checks, scorecard, invoice reconciliation, dashboard."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as CL

F = "Arial"
NAVY, TEAL, GREY = "1F3864", "2E75B6", "F2F2F2"
hfont = Font(name=F, bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", fgColor=NAVY)
cfill = PatternFill("solid", fgColor=TEAL)
base = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
GOOD = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
BAD = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")

raw = pd.read_csv("data/processed/vendor_master_raw_with_group.csv", dtype=str, keep_default_na=False)
gold = pd.read_csv("data/processed/vendor_master_golden.csv", dtype=str, keep_default_na=False)
imap = pd.read_csv("data/processed/vendor_id_map.csv", dtype=str, keep_default_na=False)
fuzzy = pd.read_csv("data/processed/duplicate_review_candidates.csv", dtype=str, keep_default_na=False)
steward = pd.read_csv("data/processed/steward_action_queue.csv", dtype=str, keep_default_na=False)
ap = pd.read_csv("data/raw/ap_invoices_raw.csv", dtype={"Invoice_Amount_INR": int})
NR, NG, NA = len(raw), len(gold), len(ap)

wb = Workbook(); wb.remove(wb.active)
def sheet(name, tab=None):
    ws = wb.create_sheet(name)
    if tab: ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    return ws
def header(ws, row, labels, fill=hfill):
    for c, t in enumerate(labels, 1):
        x = ws.cell(row, c, t); x.font, x.fill, x.border = hfont, fill, box
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
def widths(ws, w):
    for i, v in enumerate(w, 1): ws.column_dimensions[CL(i)].width = v
def put_df(ws, df, r0=1, text_cols=()):
    header(ws, r0, list(df.columns))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.font = base; c.border = box
    ws.freeze_panes = ws.cell(r0 + 1, 1)

# ======================= sheets in display order =======================
ov = sheet("Project_Overview", NAVY); db = sheet("Dashboard", "C00000"); sc = sheet("DQ_Scorecard", "C00000")
ru = sheet("DQ_Rules", TEAL); vr = sheet("Vendor_Master_Raw", "7F7F7F"); vg = sheet("Vendor_Master_Golden", "548235")
mp = sheet("Vendor_ID_Map", "548235"); fz = sheet("Duplicate_Review", "BF8F00"); sq = sheet("Steward_Queue", "BF8F00")
ia = sheet("AP_Invoices_Recon", "7030A0")

# ---------------- Vendor sheets (raw + golden) with formula checks ----------------
CHK = ["PAN_Check","GSTIN_Check","Email_Check","Phone_Check","Bank_Check","IFSC_Check","Country_Check","Terms_Check","Issue_Count","Phone_Digits_(helper)"]
def vendor_sheet(ws, df, n, with_dup):
    cols = ["Vendor_ID","Vendor_Name","PAN","GSTIN","Email","Phone","Bank_Account","IFSC","Country","Payment_Terms","Status","Created_Date"]
    header(ws, 1, cols); header(ws, 1, CHK, cfill) if False else None
    for j, t in enumerate(CHK, 13):
        x = ws.cell(1, j, t); x.font, x.fill, x.border = hfont, cfill, box
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if with_dup:
        for j, t in ((23, "Golden_Vendor_ID (from Python de-dup)"), (24, "Dup_Flag")):
            x = ws.cell(1, j, t); x.font, x.fill, x.border = hfont, cfill, box
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df[cols].itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.font, c.border = base, box
        # --- live formula checks (mirror src/dq_rules.py) ---
        ws.cell(i, 13, f'=IF(TRIM(C{i})="","Missing",IF(AND(LEN(TRIM(C{i}))=10,SUMPRODUCT(--(CODE(MID(TRIM(C{i}),{{1,2,3,4,5,10}},1))>=65),--(CODE(MID(TRIM(C{i}),{{1,2,3,4,5,10}},1))<=90))=6,ISNUMBER(--MID(TRIM(C{i}),6,4))),"Valid","Invalid"))')
        ws.cell(i, 14, f'=IF(TRIM(D{i})="","Missing",IF(AND(LEN(TRIM(D{i}))=15,MID(TRIM(D{i}),14,1)="Z"),IF(AND(TRIM(C{i})<>"",MID(TRIM(D{i}),3,10)=TRIM(C{i})),"Valid","PAN Mismatch"),"Invalid"))')
        ws.cell(i, 15, f'=IF(TRIM(E{i})="","Missing",IF(AND(ISNUMBER(FIND("@",E{i})),ISERROR(FIND(" ",TRIM(E{i}))),IFERROR(FIND(".",E{i},FIND("@",E{i})),0)>0),"Valid","Invalid"))')
        ws.cell(i, 16, f'=IF(TRIM(F{i})="","Missing",IFERROR(IF(AND(ISNUMBER(--V{i}),OR(LEN(V{i})=10,AND(LEN(V{i})=12,LEFT(V{i},2)="91")),VALUE(LEFT(RIGHT(V{i},10),1))>=6),"Valid","Invalid"),"Invalid"))')
        ws.cell(i, 17, f'=IF(TRIM(G{i})="","Missing",IFERROR(IF(AND(ISNUMBER(--TRIM(G{i})),LEN(TRIM(G{i}))>=9,LEN(TRIM(G{i}))<=18,ISERROR(FIND(".",G{i}))),"Valid","Invalid"),"Invalid"))')
        ws.cell(i, 18, f'=IF(TRIM(H{i})="","Missing",IF(AND(LEN(TRIM(H{i}))=11,SUMPRODUCT(--(CODE(MID(TRIM(H{i}),{{1,2,3,4}},1))>=65),--(CODE(MID(TRIM(H{i}),{{1,2,3,4}},1))<=90))=4,MID(TRIM(H{i}),5,1)="0"),"Valid","Invalid"))')
        ws.cell(i, 19, f'=IF(TRIM(I{i})="","Missing",IF(EXACT(TRIM(I{i}),"India"),"Valid","Invalid"))')
        ws.cell(i, 20, f'=IF(TRIM(J{i})="","Missing",IF(OR(EXACT(TRIM(J{i}),"Net 15"),EXACT(TRIM(J{i}),"Net 30"),EXACT(TRIM(J{i}),"Net 45"),EXACT(TRIM(J{i}),"Net 60")),"Valid","Invalid"))')
        ws.cell(i, 21, f'=COUNTIF(M{i}:T{i},"<>Valid")')
        ws.cell(i, 22, f'=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(F{i}," ",""),"-",""),"+","")')
        if with_dup:
            ws.cell(i, 23, df["Golden_Vendor_ID"].iloc[i - 2])
            ws.cell(i, 24, f'=IF(COUNTIF($W$2:$W${n+1},W{i})>1,"Duplicate","Unique")')
        for j in range(13, 25 if with_dup else 23):
            ws.cell(i, j).font, ws.cell(i, j).border = base, box
    last = n + 1
    rng = f"M2:T{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Valid"'], fill=GOOD))
    ws.conditional_formatting.add(rng, CellIsRule(operator="notEqual", formula=['"Valid"'], fill=BAD))
    if with_dup:
        ws.conditional_formatting.add(f"X2:X{last}", CellIsRule(operator="equal", formula=['"Duplicate"'], fill=BAD))
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{'X' if with_dup else 'V'}{last}"
    widths(ws, [10,40,13,17,34,17,18,13,9,13,10,12,11,13,10,10,10,10,12,12,10,16,16,12])
    ws.row_dimensions[1].height = 42
    for i in range(2, last + 1):
        ws.cell(i, 7).number_format = "@"
vendor_sheet(vr, raw, NR, True)
vendor_sheet(vg, gold, NG, False)
RL, GL = NR + 1, NG + 1

# ---------------- Vendor_ID_Map / review / steward ----------------
put_df(mp, imap); widths(mp, [14, 18, 24, 42, 42]); mp.auto_filter.ref = f"A1:E{len(imap)+1}"
put_df(fz, fuzzy); widths(fz, [10, 44, 13, 10, 44, 13, 11, 58]); fz.auto_filter.ref = f"A1:H{len(fuzzy)+1}"
put_df(sq, steward); widths(sq, [10, 44, 10, 62, 11, 10, 70]); sq.auto_filter.ref = f"A1:G{len(steward)+1}"
sq.conditional_formatting.add(f"F2:F{len(steward)+1}", CellIsRule(operator="equal", formula=['"High"'], fill=BAD))

# ---------------- DQ_Rules ----------------
rules = [
 ("R01","PAN","Not blank; 10 chars: 5 capital letters + 4 digits + 1 capital letter","High","Uppercase automatically; blanks need vendor document"),
 ("R02","GSTIN","15 chars, 14th char = Z, characters 3-12 must equal the vendor's PAN. Fails as 'PAN Mismatch' if PAN is missing or different","High","Cannot be auto-fixed - confirm with vendor / GST portal"),
 ("R03","Email","Not blank; contains @; a dot after the @; no spaces","Medium","Lowercase automatically; invalid needs vendor confirmation"),
 ("R04","Phone","Digits only after removing space, hyphen, +; 10 digits (or 12 starting 91); first digit 6-9","Medium","Re-formatted to +91 XXXXXXXXXX when valid"),
 ("R05","Bank_Account","Not blank; digits only; 9-18 digits","High","Cannot be auto-fixed - needs bank proof"),
 ("R06","IFSC","11 chars: 4 capital letters + '0' + 6 characters","High","Uppercase automatically; invalid needs bank proof"),
 ("R07","Country","Must equal 'India' (standard value)","Low","IN / INDIA / india mapped to India"),
 ("R08","Payment_Terms","Standard values only: Net 15 / Net 30 / Net 45 / Net 60","Low","30D, NET30, 30 Days mapped to Net 30"),
 ("R09","Duplicate vendor","Same normalised name (case, punctuation, Pvt/Private/Ltd/Limited/LLP ignored) OR same valid PAN","High","Auto-merged to one golden record; most complete record survives"),
 ("R10","Near-duplicate","Name similarity >= 0.90 but no exact name/PAN match","Medium","Listed in Duplicate_Review - NOT auto-merged"),
 ("R11","Invoice vendor","Invoice Vendor_ID must exist in vendor master","High","Orphan IDs routed to vendor onboarding"),
 ("R12","Duplicate invoice","Same golden vendor + Invoice_No + amount appears more than once","High","Hold duplicate before payment run"),
]
header(ru, 1, ["Rule_ID","Field / Area","Validation Rule","Severity","Treatment"])
for i, r in enumerate(rules, 2):
    for j, v in enumerate(r, 1):
        c = ru.cell(i, j, v); c.font, c.border = base, box; c.alignment = Alignment(wrap_text=True, vertical="top")
widths(ru, [9, 18, 80, 10, 55]); ru.freeze_panes = "A2"

# ---------------- AP_Invoices_Recon ----------------
hd = ["Invoice_No","Vendor_ID_Raw","Invoice_Date","Invoice_Amount_INR","Payment_Status","Golden_Vendor_ID","Vendor_Name (golden)","Vendor_Status",
      "Master_Data_Check","Dup_Seq_by_RawID","Dup_Seq_by_GoldenID","Dup_Flag","Missed_Without_MDM"]
header(ia, 1, hd)
for i, r in enumerate(ap.itertuples(index=False), 2):
    for j, v in enumerate([r.Invoice_No, r.Vendor_ID, r.Invoice_Date, r.Invoice_Amount_INR, r.Payment_Status], 1):
        ia.cell(i, j, v)
    ia.cell(i, 6, f'=IFERROR(INDEX(Vendor_ID_Map!$B$2:$B${len(imap)+1},MATCH(B{i},Vendor_ID_Map!$A$2:$A${len(imap)+1},0)),"UNMAPPED")')
    ia.cell(i, 7, f'=IFERROR(INDEX(Vendor_Master_Golden!$B$2:$B${GL},MATCH(F{i},Vendor_Master_Golden!$A$2:$A${GL},0)),"-")')
    ia.cell(i, 8, f'=IFERROR(INDEX(Vendor_Master_Golden!$K$2:$K${GL},MATCH(F{i},Vendor_Master_Golden!$A$2:$A${GL},0)),"-")')
    ia.cell(i, 9, f'=IF(F{i}="UNMAPPED","Orphan Vendor ID",IF(H{i}="Inactive","Inactive Vendor","OK"))')
    ia.cell(i, 10, f'=COUNTIFS($B$2:B{i},B{i},$A$2:A{i},A{i},$D$2:D{i},D{i})')
    ia.cell(i, 11, f'=COUNTIFS($F$2:F{i},F{i},$A$2:A{i},A{i},$D$2:D{i},D{i})')
    ia.cell(i, 12, f'=IF(K{i}>1,"Duplicate","Unique")')
    ia.cell(i, 13, f'=IF(AND(K{i}>1,J{i}=1),"Yes","No")')
    for j in range(1, 14): ia.cell(i, j).font, ia.cell(i, j).border = base, box
    ia.cell(i, 4).number_format = "#,##0"
IL = NA + 1
ia.conditional_formatting.add(f"L2:L{IL}", CellIsRule(operator="equal", formula=['"Duplicate"'], fill=BAD))
ia.conditional_formatting.add(f"I2:I{IL}", CellIsRule(operator="notEqual", formula=['"OK"'], fill=BAD))
ia.conditional_formatting.add(f"M2:M{IL}", CellIsRule(operator="equal", formula=['"Yes"'], fill=BAD))
widths(ia, [16,13,13,17,13,15,42,13,17,13,15,11,18]); ia.freeze_panes = "B2"; ia.auto_filter.ref = f"A1:M{IL}"; ia.row_dimensions[1].height = 32

# ---------------- DQ_Scorecard ----------------
sc["A1"] = "Vendor Master Data Quality Scorecard - Before vs After"; sc["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
sc["A2"] = "All figures are live formulas over Vendor_Master_Raw (before) and Vendor_Master_Golden (after)."; sc["A2"].font = Font(name=F, size=9, italic=True)
header(sc, 4, ["Rule","Check column","Records (Before)","Valid (Before)","Pass % (Before)","Records (After)","Valid (After)","Pass % (After)","Improvement (pts)","Open issues (After)"])
names = ["PAN","GSTIN","Email","Phone","Bank Account","IFSC","Country","Payment Terms"]
for k, nm in enumerate(names):
    r = 5 + k; col = CL(13 + k)
    sc.cell(r, 1, nm); sc.cell(r, 2, CHK[k])
    sc.cell(r, 3, f"=COUNTA(Vendor_Master_Raw!$A$2:$A${RL})")
    sc.cell(r, 4, f'=COUNTIF(Vendor_Master_Raw!${col}$2:${col}${RL},"Valid")')
    sc.cell(r, 5, f"=IF(C{r}=0,0,D{r}/C{r})")
    sc.cell(r, 6, f"=COUNTA(Vendor_Master_Golden!$A$2:$A${GL})")
    sc.cell(r, 7, f'=COUNTIF(Vendor_Master_Golden!${col}$2:${col}${GL},"Valid")')
    sc.cell(r, 8, f"=IF(F{r}=0,0,G{r}/F{r})")
    sc.cell(r, 9, f"=(H{r}-E{r})*100")
    sc.cell(r, 10, f"=F{r}-G{r}")
r = 13
sc.cell(r, 1, "Average pass %"); sc.cell(r, 5, "=AVERAGE(E5:E12)"); sc.cell(r, 8, "=AVERAGE(H5:H12)"); sc.cell(r, 9, "=(H13-E13)*100"); sc.cell(r, 10, "=SUM(J5:J12)")
sc.cell(14, 1, "Records with zero issues"); sc.cell(14, 4, f"=COUNTIF(Vendor_Master_Raw!$U$2:$U${RL},0)"); sc.cell(14, 5, "=D14/C5")
sc.cell(14, 7, f"=COUNTIF(Vendor_Master_Golden!$U$2:$U${GL},0)"); sc.cell(14, 8, "=G14/F5"); sc.cell(14, 9, "=(H14-E14)*100")
for rr in range(5, 15):
    for cc in range(1, 11):
        x = sc.cell(rr, cc); x.font = bold if rr >= 13 else base; x.border = box
        if cc in (5, 8): x.number_format = "0.0%"
        if cc == 9: x.number_format = "+0.0;-0.0;0.0"
        if rr >= 13: x.fill = PatternFill("solid", fgColor=GREY)
sc.cell(16, 1, "Note: 'After' improves on format/standardisation and duplicates. Missing or unverifiable PAN, GSTIN, bank and IFSC values are NOT invented - they go to Steward_Queue.").font = Font(name=F, size=9, italic=True)
widths(sc, [26, 20, 16, 14, 14, 16, 14, 14, 15, 16]); sc.row_dimensions[4].height = 32

# ---------------- Dashboard ----------------
db["A1"] = "Vendor Master & AP Data Governance - Dashboard"; db["A1"].font = Font(name=F, size=16, bold=True, color=NAVY)
db["A2"] = "Finance Operations | Synthetic data | All KPIs are live formulas"; db["A2"].font = Font(name=F, size=9, italic=True)
I = "AP_Invoices_Recon"
kp = [
 ("Raw vendor records", f"=COUNTA(Vendor_Master_Raw!A2:A{RL})", "#,##0"),
 ("Golden (unique) vendors", f"=COUNTA(Vendor_Master_Golden!A2:A{GL})", "#,##0"),
 ("Duplicate records merged", "=B4-B5", "#,##0"),
 ("Avg DQ pass % (before)", "=DQ_Scorecard!E13", "0.0%"),
 ("Avg DQ pass % (after)", "=DQ_Scorecard!H13", "0.0%"),
 ("Vendors needing steward action", f"=COUNTA(Steward_Queue!A2:A{len(steward)+1})", "#,##0"),
 ("High-priority steward cases", f'=COUNTIF(Steward_Queue!F2:F{len(steward)+1},"High")', "#,##0"),
 ("Near-duplicate pairs to review", f"=COUNTA(Duplicate_Review!A2:A{len(fuzzy)+1})", "#,##0"),
]
header(db, 3, ["Vendor master KPI", "Value"])
for k, (a, f, nf) in enumerate(kp, 4):
    db.cell(k, 1, a).font = base; c = db.cell(k, 2, f); c.font = bold; c.number_format = nf
    db.cell(k, 1).border = db.cell(k, 2).border = box
header(db, 13, ["AP invoice KPI", "Count", "Amount (INR)"])
ap_rows = [
 ("Total invoices", f"=COUNTA({I}!A2:A{IL})", f"=SUM({I}!D2:D{IL})"),
 ("Orphan vendor ID (not in master)", f'=COUNTIF({I}!I2:I{IL},"Orphan Vendor ID")', f'=SUMIFS({I}!D2:D{IL},{I}!I2:I{IL},"Orphan Vendor ID")'),
 ("Invoices on inactive vendors", f'=COUNTIF({I}!I2:I{IL},"Inactive Vendor")', f'=SUMIFS({I}!D2:D{IL},{I}!I2:I{IL},"Inactive Vendor")'),
 ("Duplicate invoices (after MDM)", f'=COUNTIF({I}!L2:L{IL},"Duplicate")', f'=SUMIFS({I}!D2:D{IL},{I}!L2:L{IL},"Duplicate")'),
 ("  of which duplicate already Paid", f'=COUNTIFS({I}!L2:L{IL},"Duplicate",{I}!E2:E{IL},"Paid")', f'=SUMIFS({I}!D2:D{IL},{I}!L2:L{IL},"Duplicate",{I}!E2:E{IL},"Paid")'),
 ("  of which still Pending (stoppable)", f'=COUNTIFS({I}!L2:L{IL},"Duplicate",{I}!E2:E{IL},"Pending")', f'=SUMIFS({I}!D2:D{IL},{I}!L2:L{IL},"Duplicate",{I}!E2:E{IL},"Pending")'),
 ("Duplicates found ONLY after vendor merge", f'=COUNTIF({I}!M2:M{IL},"Yes")', f'=SUMIFS({I}!D2:D{IL},{I}!M2:M{IL},"Yes")'),
]
for k, (a, c1, c2) in enumerate(ap_rows, 14):
    db.cell(k, 1, a).font = base
    db.cell(k, 2, c1).number_format = "#,##0"; db.cell(k, 3, c2).number_format = "#,##0"
    for j in (1, 2, 3):
        db.cell(k, j).border = box
        if j > 1: db.cell(k, j).font = bold
widths(db, [42, 14, 18, 3, 12, 12, 12, 12, 12, 12])
# chart 1: pass % before vs after
c1 = BarChart(); c1.type = "col"; c1.title = "DQ pass % by field: before vs after"; c1.height, c1.width = 9, 17
c1.add_data(Reference(sc, min_col=5, min_row=4, max_row=12), titles_from_data=True)
c1.add_data(Reference(sc, min_col=8, min_row=4, max_row=12), titles_from_data=True)
c1.set_categories(Reference(sc, min_col=1, min_row=5, max_row=12))
c1.y_axis.number_format = "0%"; c1.y_axis.scaling.min = 0; c1.y_axis.scaling.max = 1; c1.y_axis.delete = False; c1.x_axis.delete = False
c1.series[0].graphicalProperties.solidFill = "A6A6A6"; c1.series[1].graphicalProperties.solidFill = "2E75B6"
db.add_chart(c1, "E3")
# chart 2: AP exceptions amount
c2 = BarChart(); c2.type = "bar"; c2.title = "AP exceptions - amount (INR)"; c2.height, c2.width = 9, 17
c2.add_data(Reference(db, min_col=3, min_row=13, max_row=13 + 0), titles_from_data=False) if False else None
c2.add_data(Reference(db, min_col=3, min_row=15, max_row=20), titles_from_data=False)
c2.set_categories(Reference(db, min_col=1, min_row=15, max_row=20))
c2.legend = None; c2.y_axis.number_format = "#,##0"; c2.y_axis.delete = False; c2.x_axis.delete = False
c2.series[0].graphicalProperties.solidFill = "C00000"
db.add_chart(c2, "E22")

# ---------------- Overview ----------------
ov["A1"] = "Vendor Master Data Quality & AP Reconciliation (MDM) - Finance Operations"; ov["A1"].font = Font(name=F, size=15, bold=True, color=NAVY)
lines = [
 ("Purpose", "Clean, de-duplicate and govern vendor master data, then use the golden records to reconcile AP invoices and expose duplicate-payment risk."),
 ("Data", f"SYNTHETIC / fictional. {NR} raw vendor records, {NA} AP invoices (Jan-Sep 2025). Generated with a fixed random seed (src/01_generate_data.py)."),
 ("Method", "1) Validate every field with documented rules  2) Standardise formats  3) Detect duplicates (normalised name / PAN)  4) Build golden record by survivorship  5) Map raw IDs to golden IDs  6) Reconcile invoices  7) Score before vs after  8) Route unresolved items to data stewards."),
 ("Limits", "Validation is FORMAT-level only: no live check against Income Tax / GST portals / bank, and GSTIN check-digit is not verified. Missing values are never invented."),
 ("Formulas", "Check columns (blue headers) on both vendor sheets, the scorecard, the dashboard and the invoice reconciliation are live Excel formulas. Dup group IDs and fuzzy matches come from Python (src/02_clean_and_dedupe.py)."),
]
for k, (a, b) in enumerate(lines, 3):
    ov.cell(k, 1, a).font = bold; c = ov.cell(k, 2, b); c.font = base; c.alignment = Alignment(wrap_text=True, vertical="top")
    ov.row_dimensions[k].height = 48
header(ov, 10, ["Sheet", "What it contains"])
guide = [("Dashboard","KPIs and charts"),("DQ_Scorecard","Pass % per field, before vs after"),("DQ_Rules","Documented validation rules"),
         ("Vendor_Master_Raw","Original vendor records + formula checks + duplicate flag"),("Vendor_Master_Golden","De-duplicated, standardised golden records + checks"),
         ("Vendor_ID_Map","Raw Vendor_ID -> Golden Vendor_ID with match rule (audit trail)"),("Duplicate_Review","Near-duplicates for steward decision"),
         ("Steward_Queue","Golden records with unresolved issues, prioritised"),("AP_Invoices_Recon","Invoices mapped to golden vendors; orphan, inactive and duplicate invoice checks")]
for k, (a, b) in enumerate(guide, 11):
    ov.cell(k, 1, a).font = base; ov.cell(k, 2, b).font = base; ov.cell(k, 1).border = ov.cell(k, 2).border = box
widths(ov, [22, 120])
wb.save("excel/MDM_Vendor_DQ_Governance.xlsx")
print("saved")
